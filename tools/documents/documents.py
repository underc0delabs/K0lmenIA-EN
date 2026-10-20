#!/usr/bin/env python3
"""
Read-only survey of generated documents (PDF and Excel/CSV), to validate them in QA.

Usage (from the repo root):
    python tools/documents/documents.py --file output/.../invoice.pdf [--output ...]
    python tools/documents/documents.py --file report.xlsx

What it does:
- PDF (pypdf): extracts text per page, metadata, page count and whether it is encrypted.
- Excel (openpyxl): lists sheets, dimensions and a sample of cells; CSV: sample rows.
- Checks for **unreplaced placeholders** (leftover template variables, Lorem ipsum, TODO) — the
  most common bug in a generated document.

It does not modify the document. The document-tester agent verifies that the data/calculations are
correct and the layout is right using this input and the expected data.
"""
import argparse
import json
import re
import sys
from pathlib import Path

PLACEHOLDERS = [
    ("Template variable", re.compile(r"\{\{\s*[\w.]+\s*\}\}|\$\{[\w.]+\}|%\(?[\w]+\)?[sd]|<[\w]+>")),
    ("Lorem ipsum", re.compile(r"(?i)lorem ipsum")),
    ("Pending marker", re.compile(r"\b(TODO|FIXME|XXX|TBD|PENDING|PLACEHOLDER|CHANGEME)\b")),
    ("Sample data", re.compile(r"(?i)\b(xxx+|aaa+|test test|dummy|sample text|example text)\b")),
]


def find_placeholders(text):
    """[(type, fragment)] of markers that should not remain in a final document. Pure and testable."""
    out = []
    for kind, rx in PLACEHOLDERS:
        for m in rx.finditer(text or ""):
            out.append((kind, m.group(0)[:40]))
    return out


def read_pdf(path):
    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError("pypdf is missing: install it with `pip install pypdf`.")
    r = PdfReader(path)
    if r.is_encrypted:
        try:
            r.decrypt("")
        except Exception:
            return {"type": "pdf", "encrypted": True, "pages": len(r.pages), "text": "", "metadata": {}}
    text = []
    for pg in r.pages:
        try:
            text.append(pg.extract_text() or "")
        except Exception:
            text.append("")
    meta = {}
    try:
        meta = {k.lstrip("/"): str(v) for k, v in (r.metadata or {}).items()}
    except Exception:
        pass
    full = "\n".join(text)
    return {"type": "pdf", "encrypted": bool(r.is_encrypted), "pages": len(r.pages), "text": full,
            "metadata": meta, "placeholders": find_placeholders(full)}


def read_xlsx(path, max_rows=50):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("openpyxl is missing: install it with `pip install openpyxl`.")
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets, text = {}, []
    for sheet in wb.worksheets:
        rows = []
        for i, row in enumerate(sheet.iter_rows(values_only=True)):
            if i >= max_rows:
                break
            rows.append([("" if c is None else str(c)) for c in row])
            text += [c for c in rows[-1] if c]
        sheets[sheet.title] = {"dimension": sheet.dimensions, "sample": rows}
    full = " ".join(text)
    return {"type": "xlsx", "sheets": sheets, "placeholders": find_placeholders(full)}


def read_csv(path, max_rows=50):
    import csv
    rows, text = [], []
    with open(path, encoding="utf-8", errors="replace", newline="") as f:
        for i, row in enumerate(csv.reader(f)):
            if i >= max_rows:
                break
            rows.append(row)
            text += row
    full = " ".join(text)
    return {"type": "csv", "rows": rows, "placeholders": find_placeholders(full)}


def survey(path):
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    if ext in (".csv", ".tsv"):
        return read_csv(path)
    raise RuntimeError(f"Unsupported extension: {ext} (pdf, xlsx, csv).")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="Read-only survey of generated documents (PDF/Excel/CSV).")
    ap.add_argument("--file", required=True)
    ap.add_argument("--output", default="")
    args = ap.parse_args(argv)
    if not Path(args.file).exists():
        print(f"[documents] Not found: {args.file}", file=sys.stderr)
        return 2
    try:
        inv = survey(args.file)
    except RuntimeError as ex:
        print(f"[documents] {ex}", file=sys.stderr)
        return 2
    inv["file"] = args.file
    text = json.dumps(inv, ensure_ascii=False, indent=2)
    if args.output:
        p = Path(args.output)
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_text(text, encoding="utf-8")
        print(f"[documents] Data in {p}", file=sys.stderr)
    print(text)
    ph = inv.get("placeholders") or []
    print(f"[documents] {inv['type']} · {len(ph)} unreplaced placeholder(s).", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
