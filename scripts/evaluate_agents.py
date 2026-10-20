#!/usr/bin/env python3
"""
Evals for the k0lmenIA agents: they measure the quality of what the agents GENERATE (not just the scripts).

Each case in evals/cases/<id>/case.json says which agent is evaluated, the natural-language prompt, the
inputs that are copied into an isolated input/, and the deterministic checks on what the agent leaves behind
(files that must exist, text that must or must not appear, number of cases per type,
spreadsheet columns, valid Gherkin, Missing information IDs…). See evals/README.md.

Usage (from the repo root):
    python scripts/evaluate_agents.py --list
    python scripts/evaluate_agents.py --prepare [--case ID | --agent NAME] [--workdir DIR]
    python scripts/evaluate_agents.py --run   [--case ID | --agent NAME] [--model M] [--budget-usd 2] [--judge]
    python scripts/evaluate_agents.py --verify --workdir DIR
    python scripts/evaluate_agents.py --verify-only FOLDER --case ID

--prepare   builds an isolated workspace per case (a minimal copy of the repo: CLAUDE.md, .claude/, templates/,
             scripts/, with empty input/ and output/ + the case inputs), outside the repo (by default in the
             system temp folder), so the agent never touches your input/ or your output/.
--run     prepares, runs the agent with Claude Code in headless mode (claude -p, permissions scoped to the
             workspace, no MCP connectors, with a spending cap per case: --max-budget-usd) and verifies. SPENDS TOKENS: each case is a whole conversation.
--verify  verifies the workspaces of a run that was already prepared or run (--workdir).
--verify-only  verifies outputs that already exist in FOLDER (for example, the repo root after
             asking the agent for the case by hand). For what the agent answers in the chat, save the answer
             in FOLDER/_eval/response.txt.
--judge       (optional) in addition to the checks, asks Claude to score the case rubric. It is only
             informational: it does not change the score.

Writes evals/results/<date-time>/results.json and evals-report.html (score per case and per agent,
which check failed, cost and turns of each run, and the comparison with the previous run to detect
regressions when you change an agent or a skill). evals/results/ is not versioned.
"""
import argparse
import datetime as dt
import glob
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EVALS = ROOT / "evals"
CASES = EVALS / "cases"
RESULTS = EVALS / "results"
EVAL_DIR = "_eval"  # inside each workspace: the answer and the data of the claude run

# What is copied from the repo to the workspace (the rest, like tools/ or docs/, is not needed to evaluate).
COPY = ["CLAUDE.md", "AGENTS.md", ".claude", "templates", "scripts", "requirements.txt"]
IGNORE_COPY = shutil.ignore_patterns("__pycache__", "*.pyc", "settings.local.json", "fonts")
# Tools the agent may use in the workspace (claude -p denies everything else).
TOOLS = ["Read", "Write", "Edit", "Glob", "Grep", "Agent", "Task", "TodoWrite",
         "Bash(python:*)", "Bash(python3:*)", "Bash(git -C:*)", "Bash(ls:*)"]
PRIORITIES = ["Critical", "High", "Medium", "Low"]
CASE_STATUSES = ["N/A", "Pending", "In progress", "Passed", "Failed", "Blocked"]


# ------------------------------------------------------------------ cases
def load_cases(case=None, agent=None, folder=CASES):
    cases = []
    for file in sorted(Path(folder).glob("*/case.json")):
        data = json.loads(file.read_text(encoding="utf-8"))
        data.setdefault("id", file.parent.name)
        data["_dir"] = str(file.parent)
        if (case and data["id"] != case) or (agent and data.get("agent") != agent):
            continue
        cases.append(data)
    if case and not cases:
        raise SystemExit(f"Case {case!r} does not exist in {folder}")
    return cases


def missing(case, root=ROOT):
    """What the case requires and is not in the repo (for example, an agent that does not exist yet)."""
    return [r for r in case.get("requires", []) if not (Path(root) / r).exists()]


# ------------------------------------------------------------------ workspace
def _source(src, case):
    if src.startswith("repo:"):
        return ROOT / src[5:]
    p = EVALS / src
    return p if p.exists() else Path(case["_dir"]) / src


def prepare(case, dest):
    """Builds the isolated workspace of the case in `dest` and returns its path."""
    ws = Path(dest)
    if ws.exists():
        shutil.rmtree(ws)
    ws.mkdir(parents=True)
    for name in COPY:
        src = ROOT / name
        if src.is_dir():
            shutil.copytree(src, ws / name, ignore=IGNORE_COPY)
        elif src.exists():
            shutil.copy2(src, ws / name)
    # input/ and output/ with the same folder structure (two levels), empty.
    for base in ("input", "output"):
        for d in [p for p in (ROOT / base).rglob("*") if p.is_dir() and len(p.relative_to(ROOT).parts) <= 3] + [ROOT / base]:
            (ws / d.relative_to(ROOT)).mkdir(parents=True, exist_ok=True)
    for inp in case.get("inputs", []):
        src, dst = _source(inp["from"], case), ws / inp["to"]
        dst.parent.mkdir(parents=True, exist_ok=True)
        if src.is_dir():
            shutil.copytree(src, dst, dirs_exist_ok=True)
        else:
            shutil.copy2(src, dst)
    env = {k: str(v).replace("{workspace}", ws.as_posix()) for k, v in case.get("env", {}).items()}
    (ws / ".env").write_text("".join(f"{k}={v}\n" for k, v in env.items()), encoding="utf-8")
    (ws / EVAL_DIR).mkdir()
    (ws / EVAL_DIR / "prompt.txt").write_text(case["prompt"], encoding="utf-8")
    return ws


def claude_command(case, model=None, max_turns=None, executable="claude", budget_usd=None):
    """claude -p command line for the case: headless, JSON output, scoped permissions, spending cap and no MCP connectors."""
    cap = budget_usd or case.get("budget_usd", 2)
    cmd = [executable, "-p", case["prompt"], "--output-format", "json",
           "--permission-mode", "acceptEdits", "--allowedTools", ",".join(TOOLS),
           "--max-budget-usd", f"{float(cap):g}", "--strict-mcp-config", "--setting-sources", "project"]
    if max_turns:  # only if requested: not every Claude Code version accepts it
        cmd += ["--max-turns", str(max_turns)]
    if model:
        cmd += ["--model", model]
    return cmd


def run_claude(case, ws, model=None, max_turns=None, timeout_s=1800, budget_usd=None):
    """Runs the agent in the workspace and saves the answer and the metrics in _eval/."""
    executable = shutil.which("claude")
    if not executable:
        return {"error": "The claude command was not found (install Claude Code)."}
    start = dt.datetime.now()
    try:
        r = subprocess.run(claude_command(case, model, max_turns, executable, budget_usd), cwd=ws, capture_output=True,
                           text=True, encoding="utf-8", timeout=timeout_s)
    except subprocess.TimeoutExpired:
        return {"error": f"Timed out after {timeout_s} s."}
    (Path(ws) / EVAL_DIR / "claude.json").write_text(r.stdout or "", encoding="utf-8")
    try:
        data = json.loads(r.stdout)
    except ValueError:
        return {"error": (r.stderr or r.stdout or "claude did not return JSON").strip()[-500:]}
    (Path(ws) / EVAL_DIR / "response.txt").write_text(str(data.get("result") or ""), encoding="utf-8")
    return {"cost_usd": data.get("total_cost_usd"), "turns": data.get("num_turns"),
            "duration_s": round((data.get("duration_ms") or (dt.datetime.now() - start).total_seconds() * 1000) / 1000),
            "error": data.get("result") if data.get("is_error") else None}


def judge(case, ws, model=None):
    """Optional rubric evaluated by Claude (informational). Returns {"scores": [...]} or {"error": ...}."""
    if not case.get("rubric") or not shutil.which("claude"):
        return None
    criteria = "\n".join(f"- {c}" for c in case["rubric"])
    prompt = ("You are a QA reviewer. Evaluate what the agent left in output/, input/documentation/ and _eval/response.txt "
              f"for this request: {case['prompt']!r}. Score each criterion from 0 to 2 (0 = does not meet it, 2 = meets it well):\n"
              f"{criteria}\nAnswer ONLY with JSON: {{\"scores\": [{{\"criterion\": \"...\", \"score\": 0, \"reason\": \"...\"}}]}}")
    cmd = [shutil.which("claude"), "-p", prompt, "--output-format", "json", "--allowedTools", "Read,Glob,Grep",
           "--max-budget-usd", "1", "--strict-mcp-config"] + (["--model", model] if model else [])
    try:
        r = subprocess.run(cmd, cwd=ws, capture_output=True, text=True, encoding="utf-8", timeout=600)
        text = json.loads(r.stdout).get("result", "")
        return json.loads(re.search(r"\{.*\}", text, re.S).group(0))
    except (ValueError, AttributeError, subprocess.TimeoutExpired) as ex:
        return {"error": str(ex)}


# ------------------------------------------------------------------ reading outputs
def files(root, patterns):
    """Files in `root` that match the globs (relative to the root)."""
    patterns = [patterns] if isinstance(patterns, str) else patterns
    seen = []
    for p in patterns:
        for f in sorted(glob.glob(str(Path(root) / p), recursive=True)):
            if Path(f).is_file() and f not in seen:
                seen.append(f)
    return [Path(f) for f in seen]


def text_of(paths):
    parts = []
    for r in paths:
        if r.suffix.lower() == ".xlsx":
            parts.append("\n".join(" | ".join(str(v) for v in row.values()) for row in read_xlsx(r)[1]))
        else:
            parts.append(r.read_text(encoding="utf-8", errors="replace"))
    return "\n".join(parts)


def read_xlsx(path):
    """(headers, rows, sheets) of the test cases spreadsheet: the header row is the one with "ID"."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(f) for f in ws.iter_rows(values_only=True)]
    idx = next((i for i, f in enumerate(rows[:15]) if any(str(c or "").strip() == "ID" for c in f)), None)
    if idx is None:
        return [], [], {h.title: h for h in wb.worksheets}
    hdr = [str(c or "").strip() for c in rows[idx]]
    cases = []
    for f in rows[idx + 1:]:
        if not str(f[0] or "").strip():
            break
        cases.append({hdr[i]: ("" if v is None else str(v)) for i, v in enumerate(f) if i < len(hdr) and hdr[i]})
    return hdr, cases, {h.title: h for h in wb.worksheets}


def parse_gherkin(text):
    """Minimal structure of a .feature: {feature, feature_tags, scenarios: [...], errors: [...]}."""
    res = {"feature": None, "feature_tags": [], "scenarios": [], "errors": []}
    tags, current, in_examples = [], None, False
    for n, line in enumerate(text.splitlines(), 1):
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        if s.startswith("@"):
            tags += s.split()
            continue
        m = re.match(r"(Feature|Background|Scenario Outline|Scenario Template|Scenario|Example|Examples|Rule):\s*(.*)", s)
        if m:
            key = m.group(1)
            if key == "Feature":
                if res["feature"] is not None:
                    res["errors"].append(f"Line {n}: more than one Feature")
                res["feature"], res["feature_tags"] = m.group(2), tags
            elif key in ("Scenario", "Scenario Outline", "Scenario Template", "Example"):
                current = {"title": m.group(2), "tags": tags, "steps": [], "outline": key != "Scenario" and key != "Example",
                           "examples": 0, "line": n}
                res["scenarios"].append(current)
            elif key == "Examples":
                in_examples = True
                if current is None or not current["outline"]:
                    res["errors"].append(f"Line {n}: Examples outside a Scenario Outline")
            elif key == "Background":
                current = None
            tags = []
            if key != "Examples":
                in_examples = False
            continue
        if s.startswith("|"):
            if in_examples and current:
                current["examples"] += 1
            continue
        m = re.match(r"(Given|When|Then|And|But|\*)\s+\S", s)
        if m:
            if current is not None:
                current["steps"].append(m.group(1))
            continue
        if s.startswith('"""') or s.startswith("```"):
            continue
        if res["feature"] is None or current is not None:
            res["errors"].append(f"Line {n}: not valid Gherkin ({s[:40]})")
    if res["feature"] is None:
        res["errors"].append("It has no Feature")
    for sc in res["scenarios"]:
        if "Then" not in sc["steps"]:
            scen_err = f"The scenario on line {sc['line']} has no Then"
            res["errors"].append(scen_err)
        if not ({"Given", "When"} & set(sc["steps"])):
            res["errors"].append(f"The scenario on line {sc['line']} has neither Given nor When")
        if sc["outline"] and sc["examples"] < 2:
            res["errors"].append(f"The Scenario Outline on line {sc['line']} has no Examples table")
    return res


def _raw_tables(text):
    """Markdown table rows outside code blocks (format_tables.py was not run)."""
    outside, raw = True, 0
    for line in text.splitlines():
        if line.strip().startswith("```"):
            outside = not outside
            continue
        if outside and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", line):
            raw += 1
    return raw


# ------------------------------------------------------------------ checks
def _r(p):
    return re.compile(p, re.I | re.M | re.S)


def check(ch, root):
    """Runs a check and returns (ok, detail)."""
    kind = ch["type"]
    paths = files(root, ch.get("files", [])) if ch.get("files") else []
    if kind == "exists":
        ok = len(paths) >= ch.get("min", 1)
        return ok, f"{len(paths)} file(s): " + ", ".join(Path(r).name for r in paths[:5]) if paths else "No file was found"
    if ch.get("files") and not paths:
        return False, "The file to check was not found"
    if kind in ("contains", "not_contains"):
        text = text_of(paths)
        found = [p for p in ch["patterns"] if _r(p).search(text)]
        if kind == "not_contains":
            return not found, ("Appears: " + "; ".join(found)) if found else "Does not appear"
        if ch.get("mode") == "any":
            return bool(found), "Matches: " + "; ".join(found) if found else "No pattern matches"
        absent = [p for p in ch["patterns"] if p not in found]
        return not absent, ("Missing: " + "; ".join(absent)) if absent else "All appear"
    if kind == "count":
        matches = _r(ch["pattern"]).findall(text_of(paths))
        matches = [m if isinstance(m, str) else m[0] for m in matches]
        n = len(set(x.upper() for x in matches)) if ch.get("distinct") else len(matches)
        return n >= ch.get("min", 1), f"{n} (minimum {ch.get('min', 1)})"
    if kind == "order":
        text, pos = text_of(paths), 0
        for p in ch["patterns"]:
            m = _r(p).search(text, pos)
            if not m:
                return False, f"Does not appear (or not in order): {p}"
            pos = m.end()
        return True, "In order"
    if kind == "normalized_tables":
        raw = sum(_raw_tables(r.read_text(encoding="utf-8", errors="replace")) for r in paths)
        return raw == 0, "Aligned" if not raw else f"{raw} table(s) in raw Markdown (format_tables.py missing)"
    if kind.startswith("xlsx_"):
        return _check_xlsx(ch, paths[0])
    if kind == "gherkin":
        return _check_gherkin(ch, paths)
    raise ValueError(f"Unknown check type: {kind}")


def _case_text(c, keys=("Title", "Description", "Test Data", "#Tags")):
    return " ".join(c.get(k, "") for k in keys)


def _check_xlsx(ch, path):
    hdr, cases, sheets = read_xlsx(path)
    kind = ch["type"]
    if kind == "xlsx_columns":
        absent = [c for c in ch["columns"] if c not in hdr]
        return not absent, ("Missing columns: " + ", ".join(absent)) if absent else "All columns present"
    if kind == "xlsx_cases":
        problems = []
        if len(cases) < ch.get("min", 1):
            problems.append(f"{len(cases)} cases (minimum {ch.get('min', 1)})")
        ids = [c.get("ID", "") for c in cases]
        bad = [i for i in ids if not re.fullmatch(ch.get("ids", r"TC-\d{3}"), i)]
        if bad:
            problems.append("IDs with an invalid format: " + ", ".join(bad[:5]))
        if len(set(ids)) != len(ids):
            problems.append("Repeated IDs")
        expected = [f"TC-{i:03d}" for i in range(1, len(ids) + 1)]
        if ch.get("consecutive", True) and not bad and ids != expected:
            problems.append("IDs not consecutive from TC-001")
        prios = [c.get("Priority", "") for c in cases]
        if any(p not in PRIORITIES for p in prios):
            problems.append("Priorities outside the scale: " + ", ".join(sorted({p for p in prios if p not in PRIORITIES})))
        elif ch.get("priority_order", True) and [PRIORITIES.index(p) for p in prios] != sorted(PRIORITIES.index(p) for p in prios):
            problems.append("Not sorted by priority (Critical → Low)")
        statuses = {c.get("Status", "") for c in cases} - set(CASE_STATUSES)
        if statuses:
            problems.append("Statuses outside the scale: " + ", ".join(sorted(statuses)))
        no_result = [c.get("ID") for c in cases if not c.get("Expected Result", "").strip() or not c.get("Steps", "").strip()]
        if no_result:
            problems.append("No steps or expected result: " + ", ".join(no_result[:5]))
        return not problems, "; ".join(problems) or f"{len(cases)} valid cases"
    if kind == "xlsx_types":
        detail, ok = [], True
        for name, rule in ch["types"].items():
            n = sum(1 for c in cases if _r(rule["pattern"]).search(_case_text(c)))
            ok &= n >= rule["min"]
            detail.append(f"{name}: {n} (min. {rule['min']})")
        return ok, " · ".join(detail)
    if kind == "xlsx_missing_info":
        marked = [c for c in cases if "@missing-info" in c.get("#Tags", "") or re.search(r"MI-\d+", c.get("Comments", ""))]
        sheet = sheets.get("Missing information")
        items = set()
        if sheet is not None:
            for row in sheet.iter_rows(values_only=True):
                items |= {m for v in row for m in re.findall(r"MI-\d+", str(v or ""))}
        ok = len(marked) >= ch.get("min_cases", 1) and len(items) >= ch.get("min_items", 1)
        return ok, f"{len(marked)} case(s) marked · {len(items)} item(s) in the Missing information sheet"
    if kind == "xlsx_filter":
        affected = [c for c in cases if _r(ch["filter"]).search(_case_text(c, ("Title", "Description", "Steps", "Expected Result")))]
        if not affected:
            return ch.get("ok_if_none", False), "No case covers the topic"
        bad = []
        for c in affected:
            marks = c.get("#Tags", "") + " " + c.get("Comments", "")
            if ch.get("must") and not _r(ch["must"]).search(marks):
                bad.append(f"{c.get('ID')} without {ch['must']}")
            if ch.get("must_not_result") and _r(ch["must_not_result"]).search(c.get("Expected Result", "")):
                bad.append(f"{c.get('ID')} invents: {c.get('Expected Result', '')[:60]}")
        return not bad, "; ".join(bad) or f"{len(affected)} correct case(s)"
    raise ValueError(f"Unknown check type: {kind}")


def _check_gherkin(ch, paths):
    problems, scenarios, ids, with_mi = [], 0, [], 0
    for r in paths:
        text = r.read_text(encoding="utf-8", errors="replace")
        g = parse_gherkin(text)
        problems += [f"{r.name}: {x}" for x in g["errors"]]
        if ch.get("tag_feature") and not any(re.fullmatch(ch["tag_feature"], t) for t in g["feature_tags"]):
            problems.append(f"{r.name}: the Feature does not have the tag {ch['tag_feature']}")
        for sc in g["scenarios"]:
            scenarios += 1
            tag_tc = [t for t in sc["tags"] if re.fullmatch(ch.get("tag_scenario", r"@TC-\d{3}"), t)]
            if not tag_tc:
                problems.append(f"{r.name}: '{sc['title'][:40]}' without a case tag")
            ids += tag_tc
            if "@missing-info" in sc["tags"] and any(re.fullmatch(r"@MI-\d+", t) for t in sc["tags"]):
                with_mi += 1
        if ch.get("missing_info_min") and not re.search(r"#\s*MISSING INFO(RMATION)? \(MI-\d+\)", text, re.I):
            problems.append(f"{r.name}: the # MISSING INFORMATION (MI-XX) comment is missing")
    if len(set(ids)) != len(ids):
        problems.append("Repeated case tags")
    if scenarios < ch.get("min_scenarios", 1):
        problems.append(f"{scenarios} scenarios (minimum {ch.get('min_scenarios', 1)})")
    if with_mi < ch.get("missing_info_min", 0):
        problems.append(f"{with_mi} scenario(s) with @missing-info @MI-XX (minimum {ch['missing_info_min']})")
    return not problems, "; ".join(problems[:6]) or f"{scenarios} valid scenarios"


def verify_case(case, root):
    """Runs all the checks of the case on `root` and returns the result with its score."""
    results, total, achieved = [], 0, 0
    for ch in case.get("checks", []):
        try:
            ok, detail = check(ch, root)
        except Exception as ex:  # a broken check does not stop the evaluation
            ok, detail = False, f"Error while verifying: {ex}"
        weight = ch.get("weight", 1)
        total += weight
        achieved += weight if ok else 0
        results.append({"id": ch["id"], "description": ch.get("description", ch["id"]), "ok": ok, "detail": detail, "weight": weight})
    score = round(100 * achieved / total) if total else 0
    status = "Passed" if score >= case.get("threshold", 80) else "Failed"
    return {"id": case["id"], "agent": case.get("agent", ""), "description": case.get("description", ""),
            "score": score, "status": status, "checks": results}


# ------------------------------------------------------------------ comparison and report
def previous_run(dest):
    previous = sorted(p for p in Path(dest).parent.glob("*/results.json") if p.parent != Path(dest))
    previous = [p for p in previous if p.parent.name < Path(dest).name]
    return json.loads(previous[-1].read_text(encoding="utf-8")) if previous else None


def compare(current, previous):
    """Per case: score before/after, checks that stopped passing (regressions) and those that now pass."""
    if not previous:
        return None
    prev = {c["id"]: c for c in previous.get("cases", [])}
    out = {"previous_date": previous.get("date"), "cases": {}, "regressions": 0, "improvements": 0}
    for c in current["cases"]:
        p = prev.get(c["id"])
        if not p or c["status"] in ("N/A", "Blocked") or p.get("status") in ("N/A", "Blocked"):
            continue
        before = {x["id"]: x["ok"] for x in p.get("checks", [])}
        reg = [x["id"] for x in c["checks"] if before.get(x["id"]) is True and not x["ok"]]
        imp = [x["id"] for x in c["checks"] if before.get(x["id"]) is False and x["ok"]]
        out["cases"][c["id"]] = {"before": p.get("score"), "after": c["score"], "regressions": reg, "improvements": imp}
        out["regressions"] += len(reg)
        out["improvements"] += len(imp)
    return out


def summary_by_agent(cases):
    agents = {}
    for c in cases:
        if c["status"] in ("N/A", "Blocked"):
            continue
        a = agents.setdefault(c["agent"], {"cases": 0, "passed": 0, "scores": []})
        a["cases"] += 1
        a["passed"] += c["status"] == "Passed"
        a["scores"].append(c["score"])
    return {k: {"cases": v["cases"], "passed": v["passed"], "score": round(sum(v["scores"]) / len(v["scores"]))}
            for k, v in agents.items()}


def html_report(data):
    sys.path.insert(0, str(ROOT / "scripts"))
    from _report_styles import (e, page, meta, section, kpi_tiles, banner, table, bar_chart, status_pill, doc_list,
                                  OK, WARNING, CRITICAL)
    cases = data["cases"]
    evaluated = [c for c in cases if c["status"] in ("Passed", "Failed")]
    passed = sum(c["status"] == "Passed" for c in evaluated)
    avg = round(sum(c["score"] for c in evaluated) / len(evaluated)) if evaluated else 0
    comp = data.get("comparison") or {}
    cost = sum((c.get("metrics") or {}).get("cost_usd") or 0 for c in cases)
    tone = "ok" if evaluated and passed == len(evaluated) and not comp.get("regressions") else "warn" if avg >= 60 else "bad"
    verdict = (f"{passed} of {len(evaluated)} cases passed · average score {avg}"
               + (f" · {comp['regressions']} regression(s) against the previous run" if comp.get("regressions") else ""))
    body = section("Result", banner(verdict, tone) + kpi_tiles([
        ("Cases evaluated", len(evaluated), ""), ("Passed", passed, "ok" if passed == len(evaluated) else "warn"),
        ("Average score", avg, "ok" if avg >= 80 else "warn" if avg >= 60 else "bad", "", "target", {"pct": avg}),
        ("Regressions", comp.get("regressions", 0) if comp else "—", "bad" if comp.get("regressions") else ""),
        ("Total cost (USD)", f"{cost:.2f}" if cost else "—", ""),
    ]))
    by_agent = data.get("by_agent", {})
    rows = [[f"<code>{e(a)}</code>", v["cases"], v["passed"], v["score"]] for a, v in sorted(by_agent.items())]
    body += section("By agent", table(["Agent", "Cases", "Passed", "Average score"], rows), width=6)
    color = lambda p: OK if p >= 80 else WARNING if p >= 60 else CRITICAL
    body += section("Score per case", bar_chart([(c["id"], c["score"], color(c["score"])) for c in evaluated], 220), width=6)
    if comp and comp.get("cases"):
        rows = [[e(cid), v["before"], v["after"], e(", ".join(v["regressions"]) or "—"), e(", ".join(v["improvements"]) or "—")]
                for cid, v in comp["cases"].items()]
        body += section("Comparison with the previous run", table(["Case", "Before", "Now", "Regressions", "Improvements"], rows),
                        sub=f"Previous run: {comp.get('previous_date', '')}")
    for c in cases:
        m = c.get("metrics") or {}
        info = " · ".join(x for x in [
            f"Agent {c['agent']}", f"Score {c['score']}" if c["status"] in ("Passed", "Failed") else "",
            f"{m['turns']} turns" if m.get("turns") else "", f"US$ {m['cost_usd']:.2f}" if m.get("cost_usd") else "",
            f"{m['duration_s']} s" if m.get("duration_s") else ""] if x)
        case_body = f'<div class="cardbody"><p class="muted">{e(c.get("description", ""))}</p><p>{status_pill(c["status"])} {e(info)}</p></div>'
        if c.get("reason"):
            case_body += doc_list([c["reason"]])
        if c.get("checks"):
            case_body += table(["Status", "Check", "Detail"], [
                [status_pill("Passed" if x["ok"] else "Failed"), e(x["description"]), e(x["detail"])] for x in c["checks"]])
        if c.get("judge") and c["judge"].get("scores"):
            case_body += table(["Criterion (judge, informational)", "Score", "Reason"], [
                [e(p.get("criterion")), e(p.get("score")), e(p.get("reason"))] for p in c["judge"]["scores"]])
        body += section(f"Case {c['id']}", case_body)
    return page("Agent evals", "Evals", meta([("Date", data["date"]), ("Model", data.get("model") or "default"),
                                               ("Cases", len(cases))]), body)


def save(data, dest):
    dest = Path(dest)
    dest.mkdir(parents=True, exist_ok=True)
    data["by_agent"] = summary_by_agent(data["cases"])
    data["comparison"] = compare(data, previous_run(dest))
    (dest / "results.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    (dest / "evals-report.html").write_text(html_report(data), encoding="utf-8")
    return dest


# ------------------------------------------------------------------ CLI
def _stamp():
    return dt.datetime.now().strftime("%Y%m%d-%H%M%S")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Evals for the k0lmenIA agents")
    action = ap.add_mutually_exclusive_group(required=True)
    action.add_argument("--list", action="store_true")
    action.add_argument("--prepare", action="store_true")
    action.add_argument("--run", action="store_true")
    action.add_argument("--verify", action="store_true")
    action.add_argument("--verify-only", metavar="FOLDER")
    ap.add_argument("--case")
    ap.add_argument("--agent")
    ap.add_argument("--workdir", help="folder for the workspaces (by default, in the system temp folder)")
    ap.add_argument("--results", default=str(RESULTS))
    ap.add_argument("--cases", default=str(CASES), help=argparse.SUPPRESS)
    ap.add_argument("--model")
    ap.add_argument("--max-turns", type=int, help="turn cap (--max-turns), if your Claude Code version accepts it")
    ap.add_argument("--budget-usd", type=float, help="spending cap per case (overrides the case budget_usd; 2 by default)")
    ap.add_argument("--judge", action="store_true")
    a = ap.parse_args(argv)
    cases = load_cases(a.case, a.agent, a.cases)
    stamp = _stamp()

    if a.list:
        for c in cases:
            absent = missing(c)
            print(f"{c['id']:32} {c.get('agent', ''):26} {len(c.get('checks', []))} checks"
                  + (f"  (skipped: missing {', '.join(absent)})" if absent else ""))
        return 0

    if a.verify_only:
        data = {"date": stamp, "mode": "verify-only", "folder": a.verify_only,
                "cases": [verify_case(c, a.verify_only) for c in cases]}
        dest = save(data, Path(a.results) / stamp)
        _print(data, dest)
        return 0 if all(c["status"] == "Passed" for c in data["cases"]) else 1

    workdir = Path(a.workdir) if a.workdir else Path(tempfile.gettempdir()) / "k0lmenia-evals" / stamp
    if a.verify:
        if not a.workdir:
            raise SystemExit("--verify needs --workdir with the run folder")
        run = json.loads((workdir / "run.json").read_text(encoding="utf-8")) if (workdir / "run.json").exists() else {}
        data = {"date": stamp, "mode": "verify", "workdir": str(workdir), "model": run.get("model"), "cases": []}
        for c in cases:
            ws = workdir / c["id"]
            if not ws.exists():
                continue
            res = verify_case(c, ws)
            res["metrics"] = (run.get("metrics") or {}).get(c["id"])
            data["cases"].append(res)
        dest = save(data, Path(a.results) / stamp)
        _print(data, dest)
        return 0

    run = {"date": stamp, "model": a.model, "cases": [], "metrics": {}}
    data = {"date": stamp, "mode": "run" if a.run else "prepare", "workdir": str(workdir), "model": a.model, "cases": []}
    for c in cases:
        absent = missing(c)
        if absent:
            data["cases"].append({"id": c["id"], "agent": c.get("agent", ""), "description": c.get("description", ""),
                                  "score": 0, "status": "N/A", "checks": [], "reason": "Skipped: missing " + ", ".join(absent)})
            continue
        ws = prepare(c, workdir / c["id"])
        run["cases"].append(c["id"])
        if a.prepare:
            print(f"{c['id']}: {ws}")
            print("   cd \"" + str(ws) + "\" && " + " ".join(_quote(x) for x in claude_command(c, a.model, a.max_turns, budget_usd=a.budget_usd)))
            continue
        print(f"Running {c['id']} ({c.get('agent')})…", flush=True)
        metrics = run_claude(c, ws, a.model, a.max_turns, budget_usd=a.budget_usd)
        run["metrics"][c["id"]] = metrics
        if metrics.get("error") and not any((ws / "output").rglob("*.*")):
            data["cases"].append({"id": c["id"], "agent": c.get("agent", ""), "description": c.get("description", ""),
                                  "score": 0, "status": "Blocked", "checks": [], "metrics": metrics,
                                  "reason": f"Could not run: {metrics['error']}"})
            continue
        res = verify_case(c, ws)
        res["metrics"] = metrics
        if a.judge:
            res["judge"] = judge(c, ws, a.model)
        data["cases"].append(res)
    workdir.mkdir(parents=True, exist_ok=True)
    (workdir / "run.json").write_text(json.dumps(run, ensure_ascii=False, indent=2), encoding="utf-8")
    if a.prepare:
        print(f"\nWorkspaces in {workdir}. After running the prompts: python scripts/evaluate_agents.py --verify --workdir \"{workdir}\"")
        return 0
    dest = save(data, Path(a.results) / stamp)
    _print(data, dest)
    return 0


def _quote(x):
    return x if re.fullmatch(r"[\w@%+=:,./*()-]+", x) else '"' + x.replace('"', '\\"') + '"'


def _print(data, dest):
    for c in data["cases"]:
        print(f"  {c['status']:10} {c['score']:>3}  {c['id']}  ({c['agent']})")
        for x in c.get("checks", []):
            if not x["ok"]:
                print(f"             ✕ {x['description']}: {x['detail']}")
    comp = data.get("comparison")
    if comp:
        print(f"Against the previous run: {comp['regressions']} regression(s), {comp['improvements']} improvement(s).")
    print(f"Report: {Path(dest) / 'evals-report.html'}")


if __name__ == "__main__":
    sys.exit(main())
