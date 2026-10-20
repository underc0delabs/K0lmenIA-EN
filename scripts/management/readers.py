"""Reading of local inputs, independent of the test management tool.

- Manual cases from the .xlsx spreadsheet of generate_cases.py.
- BDD scenarios from a .feature.
- Run results from the k0lmena cucumber-report.json (web, api or mobile),
  with their evidence (embedded screenshots and videos).
"""
import base64
import hashlib
import re
from pathlib import Path

from common import ManagementError, need, normalize_status

PRIORITIES = {"critical": "Critical", "high": "High", "medium": "Medium", "low": "Low"}


def priority_from_tags(tags):
    """The priority of a BDD scenario from its tags: an explicit tag (@critical/@high/@medium/@low)
    wins; if there is none, @smoke means High; otherwise, Medium."""
    for t in tags:
        p = PRIORITIES.get(str(t).lower())
        if p:
            return p
    return "High" if any(str(t).lower() == "smoke" for t in tags) else "Medium"


def _steps_from_text(steps, data, result):
    """'1. Go to...\\n2. Fill in...' -> [{action, data, result}]. The expected result of the
    case goes in the last step and the test data in the first one."""
    lines = [re.sub(r"^\s*\d+[.)-]\s*", "", l).strip() for l in (steps or "").splitlines()]
    lines = [l for l in lines if l]
    if not lines:
        lines = ["(no steps)"]
    out = [{"action": l, "data": "", "result": ""} for l in lines]
    out[0]["data"] = data or ""
    out[-1]["result"] = result or ""
    return out


def read_xlsx(path):
    """Reads the cases spreadsheet (format of scripts/generate_cases.py)."""
    need("openpyxl")
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True).active
    headers, header_row = None, None
    for row in ws.iter_rows(min_row=1, max_row=15):
        values = [str(c.value or "").strip() for c in row]
        if "ID" in values and "Title" in values:
            headers, header_row = values, row[0].row
            break
    if not headers:
        raise ManagementError(f"{path}: cannot find the header row (ID, Title, ...).")
    col = {name: i for i, name in enumerate(headers)}

    def val(row, name):
        i = col.get(name)
        return str(row[i] or "").strip() if i is not None and i < len(row) else ""

    cases = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        cid = val(row, "ID")
        if not re.match(r"^TC(-API)?-\d+", cid):
            continue
        cases.append({
            "id": cid,
            "title": val(row, "Title"),
            "description": val(row, "Description"),
            "preconditions": val(row, "Preconditions"),
            "priority": PRIORITIES.get(val(row, "Priority").lower(), "Medium"),
            "tags": [e.lstrip("@") for e in re.split(r"[\s,]+", val(row, "#Tags")) if e],
            "type": "manual",
            "time_min": val(row, "Estimated Time (min)"),
            "steps": _steps_from_text(val(row, "Steps"), val(row, "Test Data"), val(row, "Expected Result")),
        })
    if not cases:
        raise ManagementError(f"{path}: found no cases with ID TC-XXX.")
    return cases


_SCENARIO_START = re.compile(r"^\s*(Scenario Outline|Scenario Template|Scenario|Example|Rule):")
_BLOCK_START = re.compile(r"^\s*(Scenario Outline|Scenario Template|Scenario|Example|Rule|Background|Feature):")


def _block_end(lines, j):
    """True if another scenario, rule or feature starts on line j. A tag (@x) closes the block
    only if what follows (after more tags and comments) is a scenario or a rule: the tags
    of an "Examples:" are part of the Scenario Outline."""
    l = lines[j].strip()
    if _BLOCK_START.match(l):
        return True
    if l.startswith("@"):
        k = j + 1
        while k < len(lines) and (not lines[k].strip() or lines[k].strip().startswith(("@", "#"))):
            k += 1
        return k < len(lines) and bool(_BLOCK_START.match(lines[k].strip()))
    return False


def read_feature(path):
    """Scenarios of a .feature as cases of type 'cucumber' (the scenario's full gherkin).
    The case ID comes from the scenario's @TC-XXX tag; if it has none, the title is used.
    Supports Background (of the feature and of each Rule), Scenario Outline with several Examples
    (also with tags), Example/Scenario Template and Rule."""
    text = Path(path).read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    background, rule_background, cases, feature_tags, rule_tags = [], [], [], [], []
    i = 0
    pending_tags = []
    in_rule = False
    while i < len(lines):
        l = lines[i].strip()
        if l.startswith("@"):
            pending_tags += [t.lstrip("@") for t in l.split() if t.startswith("@")]
        elif l.startswith("Feature:"):
            feature_tags, pending_tags = pending_tags, []
        elif l.startswith("Rule:"):
            in_rule, rule_tags, pending_tags, rule_background = True, pending_tags, [], []
        elif l.startswith("Background:"):
            target = rule_background if in_rule else background
            j = i + 1
            while j < len(lines) and not _block_end(lines, j):
                target.append(lines[j].rstrip())
                j += 1
            i = j
            continue
        elif _SCENARIO_START.match(l):
            title = l.split(":", 1)[1].strip()
            body = [l]
            j = i + 1
            while j < len(lines) and not _block_end(lines, j):
                body.append(lines[j].rstrip())
                j += 1
            while body and not body[-1].strip():
                body.pop()
            tags = pending_tags
            cid = next((t for t in tags if re.match(r"^TC(-API)?-\d+$", t)), title)
            bg = [x for x in background + rule_background if x.strip()]
            gherkin = "\n".join((["Background:"] + bg + [""] if bg else []) + body)
            all_tags = feature_tags + rule_tags + tags
            cases.append({
                "id": cid, "title": title, "description": "", "preconditions": "",
                "priority": priority_from_tags(all_tags),
                "tags": list(dict.fromkeys(t for t in all_tags if not re.match(r"^TC(-API)?-\d+$", t))),
                "type": "cucumber", "gherkin": gherkin.strip(), "steps": [],
            })
            pending_tags = []
            i = j
            continue
        i += 1
    if not cases:
        raise ManagementError(f"{path}: found no scenarios.")
    return cases


def read_cases(source, ids=None):
    suffix = Path(source).suffix.lower()
    if suffix == ".xlsx":
        cases = read_xlsx(source)
    elif suffix == ".feature":
        cases = read_feature(source)
    else:
        raise ManagementError("The source of the cases must be an .xlsx (manual) or a .feature (BDD).")
    if ids:
        requested = set(ids)
        cases = [c for c in cases if c["id"] in requested]
        missing = requested - {c["id"] for c in cases}
        if missing:
            raise ManagementError(f"Could not find these cases in {source}: {', '.join(sorted(missing))}.")
    # Estimated manual run time (scripts/estimation.py): it travels to the test management tool.
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from estimation import estimate_case
    for c in cases:
        c["estimate"] = estimate_case(c)
    return cases


def read_results(report_json, evidence_dir):
    """k0lmena cucumber-report.json -> list of results with the evidence in files.

    Embedded screenshots (image/png) are written to evidence_dir. The video is taken from the
    HTML attachment (<video src="videos/...">) left by the hook. With retries, the last attempt is kept.
    """
    import json

    report = Path(report_json)
    if not report.exists():
        raise ManagementError(f"{report} does not exist. Run npm test in tools/k0lmena first.")
    report_base = report.parent
    evidence_dir = Path(evidence_dir)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    features = json.loads(report.read_text(encoding="utf-8"))

    by_scenario = {}
    for feature in features:
        for sc in feature.get("elements", []):
            if sc.get("type") == "background":
                continue
            tags = list(dict.fromkeys(t["name"].lstrip("@") for t in sc.get("tags", [])))
            statuses, error, screenshots, videos, gifs, texts = [], "", [], [], [], []
            for step in sc.get("steps", []):
                res = step.get("result", {})
                statuses.append(res.get("status", "undefined"))
                if res.get("status") == "failed" and not error:
                    error = (res.get("error_message") or "").strip()
                for emb in step.get("embeddings", []):
                    mime = emb.get("mime_type", "")
                    if mime == "image/png":
                        screenshots.append(base64.b64decode(emb["data"]))
                    elif mime == "text/html":
                        html = base64.b64decode(emb["data"]).decode("utf-8", "replace")
                        videos += re.findall(r'<video[^>]*src="([^"]+)"', html)
                        gifs += re.findall(r'<img[^>]*src="([^"]+[.]gif)"', html)
                    elif mime == "text/plain":
                        texts.append(base64.b64decode(emb["data"]).decode("utf-8", "replace"))
            if "failed" in statuses or "ambiguous" in statuses:
                status = "failed"
            elif all(s == "passed" for s in statuses):
                status = "passed"
            else:
                status = "pending"
            key = f'{feature.get("uri", "")}:{sc.get("line")}:{sc.get("name")}'
            by_scenario[key] = {
                "scenario": sc.get("name", ""), "feature": feature.get("name", ""),
                "tags": tags, "status": normalize_status(status), "error": error,
                "_screenshots": screenshots, "_videos": videos, "_gifs": gifs, "texts": texts,
            }

    results = []
    for r in by_scenario.values():
        name = re.sub(r"[^\w.-]+", "_", r["scenario"])[:80]
        evidence = []
        for n, png in enumerate(r.pop("_screenshots")):
            file = evidence_dir / f"{name}-{hashlib.md5(png).hexdigest()[:8]}-{n}.png"
            file.write_bytes(png)
            evidence.append({"type": "screenshot", "file": str(file)})
        for kind, key in (("video", "_videos"), ("gif", "_gifs")):
            for src in r.pop(key):
                # web: "videos/x.webm" (relative to reports/web) · mobile: "../videos/x.mp4" (relative to html/features)
                clean = src.replace("../", "")
                for candidate in (report_base / clean, report_base / "html" / clean):
                    if candidate.exists():
                        evidence.append({"type": kind, "file": str(candidate.resolve())})
                        break
        r["evidence"] = evidence
        results.append(r)
    return results
