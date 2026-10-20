#!/usr/bin/env python3
"""Generates the test cases from a JSON: the Excel spreadsheet (.xlsx) and,
in the same step, a Markdown version (.md) with the cases in an aligned ASCII table.

Usage:
    python scripts/generate_cases.py <input.json> <output.xlsx> [--clean]

Writes <output.xlsx> and, next to it, <output.md> (same path, .md extension).
With --clean, deletes the input JSON when done (useful for temporary files).

Input JSON format:
{
  "project": "PROJECT NAME",
  "module": "US-001 — Login",
  "cases": [
    {"id": "TC-001", "title": "...", "description": "...", "preconditions": "...",
     "data": "Email: ...\nPassword: ...", "steps": "1. ...\n2. ...", "result": "...",
     "status": "Pending", "priority": "High", "tags": "@login", "time_min": 5.5,
     "evidence": "", "run_date": "", "comments": "",
     "missing_info": ["MI-01"]}
  ],
  "missing_info": [
    {"id": "MI-01", "what_is_missing": "Text of the message when the account gets locked",
     "where_searched": "Story, comments, Figma", "impact": "AC3",
     "question": "What is the exact message?", "status": "Open"}
  ]
}

"time_min" (optional, per case): minutes of manual run time. If missing, it is calculated with the formula of
scripts/estimation.py ("Estimated Time (min)" column of the .xlsx, with the total at the bottom).
"missing_info" (optional, per case): IDs of the "Missing information" items the case depends on.
Those cases are highlighted in the spreadsheet, get the @missing-info tag and the detail in
Comments. The top-level "missing_info" (optional) builds the "Missing information" sheet of the .xlsx
and the section of the same name in the .md (see skill context-research).
"""
import sys
import os
import re
import json
import textwrap
import subprocess

COLS = ["ID", "Title", "Description", "Preconditions", "Test Data", "Steps",
        "Expected Result", "Status", "Priority", "Estimated Time (min)", "#Tags", "Evidence",
        "Run Date", "Comments"]
KEYS = ["id", "title", "description", "preconditions", "data", "steps", "result",
        "status", "priority", "time_min", "tags", "evidence", "run_date", "comments"]
# The .md has: (1) a Summary table and (2) the table with all the cases (one row per case).
# Summary: a few columns, sorted by priority (most critical first).
SUMMARY_COLS = ["ID", "Title", "Expected result", "Priority", "Time"]
SUMMARY_KEYS = ["id", "title", "result", "priority", "time"]
SUMMARY_W = [6, 24, 32, 10, 8]

# Cases table of the .md. Drops Status (always "Pending"), Evidence, Run Date
# and Comments so it does not get too wide; all those columns are in the .xlsx.
MD_COLS = ["ID", "Title", "Description", "Preconditions", "Test Data", "Steps",
           "Expected Result", "Priority", "Time", "#Tags"]
MD_KEYS = ["id", "title", "description", "preconditions", "data", "steps",
           "result", "priority", "time", "tags"]
# Maximum width per column (long cells wrap onto several lines inside the cell).
MD_MAXW = [6, 16, 22, 16, 18, 24, 20, 9, 7, 12]

MI_COLS = ["ID", "What is missing", "Where searched", "Impact", "Question for the PO", "Status"]
MI_KEYS = ["id", "what_is_missing", "where_searched", "impact", "question", "status"]
MI_W = [8, 40, 30, 22, 40, 11]


def _mi_ids(case):
    v = case.get("missing_info") or []
    if isinstance(v, str):
        v = [x for x in v.replace(";", ",").split(",")]
    return [str(x).strip() for x in v if str(x).strip()]


def prepare_missing_info(data):
    """Flags the cases that depend on missing information: @missing-info tag and
    the detail of each MI in Comments (without duplicating it if it was already there)."""
    detail = {str(f.get("id", "")).strip(): f for f in data.get("missing_info", []) or []}
    for case in data.get("cases", []):
        ids = _mi_ids(case)
        if not ids:
            continue
        tags = str(case.get("tags", "") or "")
        if "@missing-info" not in tags:
            case["tags"] = (tags + " @missing-info").strip()
        comments = str(case.get("comments", "") or "")
        notes = []
        for mid in ids:
            if mid in comments:
                continue
            what = (detail.get(mid) or {}).get("what_is_missing", "")
            notes.append(f"MISSING INFORMATION ({mid}){': ' + what if what else ''}")
        if notes:
            case["comments"] = "\n".join(notes + ([comments] if comments else []))
        case["_mi"] = ", ".join(ids)


def prepare_times(data):
    """Estimated manual run time of each case (scripts/estimation.py). If the JSON already has
    "time_min" (someone adjusted it), it is kept."""
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from estimation import estimate_case, format_time
    total = 0.0
    for case in data.get("cases", []):
        try:
            minutes = float(str(case.get("time_min")).replace(",", "."))
        except (TypeError, ValueError):
            minutes = estimate_case(case)["minutes"]
        case["time_min"] = minutes
        case["time"] = format_time(minutes)
        total += minutes
    data["_total_time"] = total
    data["_total_time_txt"] = format_time(total)


# Priority order: most critical first, low priority last.
PRIORITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3}


def _need(pkg, install=None):
    """Makes sure the library is available. If it is missing, tries to install it; if it
    cannot, warns with a clear message and stops (instead of failing with a traceback)."""
    install = install or pkg
    try:
        __import__(pkg)
        return
    except ImportError:
        pass
    subprocess.run([sys.executable, "-m", "pip", "install", install, "--quiet"], check=False)
    try:
        __import__(pkg)
    except ImportError:
        sys.exit(
            f"\n[!] The library '{install}' is missing and could not be installed automatically.\n"
            f"    Install the dependencies with:  pip install -r requirements.txt\n"
            f"    (or directly:  pip install {install})\n")


def to_xlsx(data, out):
    _need("openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    project = data.get("project", "PROJECT NAME")
    module = data.get("module", "Module")
    cases = data.get("cases", [])
    widths = [6, 30, 42, 32, 30, 46, 40, 13, 11, 12, 22, 20, 18, 30]
    n = len(COLS)
    last_col = chr(ord("A") + n - 1)

    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[:\\/?*\[\]]", "-", module)[:31]) or "Cases"
    gray = PatternFill("solid", fgColor="404040")
    blue = PatternFill("solid", fgColor="4472C4")
    white_bold = Font(name="Arial", color="FFFFFF", bold=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = project
    ws["A1"].fill = gray
    ws["A1"].font = Font(name="Arial", color="FFFFFF", bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = f"Module: {module}"
    ws["A2"].fill = blue
    ws["A2"].font = white_bold
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    for i, c in enumerate(COLS):
        cell = ws.cell(row=3, column=i + 1, value=c)
        cell.fill = blue
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[3].height = 24
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = w

    data_font = Font(name="Arial", size=10)
    amber = PatternFill("solid", fgColor="FFF2CC")
    amber_strong = Font(name="Arial", size=10, bold=True, color="9C5700")
    start = 4
    for r, case in enumerate(cases):
        row = start + r
        missing = bool(case.get("_mi"))
        for c, key in enumerate(KEYS):
            cell = ws.cell(row=row, column=c + 1, value=case.get(key, ""))
            cell.font = amber_strong if (missing and key == "comments") else data_font
            cell.border = border
            if missing:
                cell.fill = amber
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if c in (0, 7, 8, 9) else "left")
        ws.row_dimensions[row].height = 95

    # Total of the manual run (the time column is numeric: it can be summed or filtered in Excel).
    total_row = start + len(cases)
    col_t = KEYS.index("time_min") + 1
    letter_t = chr(ord("A") + col_t - 1)
    et = ws.cell(row=total_row, column=col_t - 1, value="Estimated total")
    et.font = Font(name="Arial", size=10, bold=True)
    et.alignment = Alignment(horizontal="right")
    tot = ws.cell(row=total_row, column=col_t, value=f"=SUM({letter_t}{start}:{letter_t}{total_row - 1})" if cases else 0)
    tot.font = Font(name="Arial", size=10, bold=True)
    tot.alignment = Alignment(horizontal="center")
    tot.border = border

    last_row = start + len(cases) + 50
    dv_status = DataValidation(type="list",
        formula1='"N/A,Pending,In progress,Passed,Failed,Blocked"', allow_blank=True)
    dv_prio = DataValidation(type="list", formula1='"Critical,High,Medium,Low"', allow_blank=True)
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_prio)
    dv_status.add(f"H{start}:H{last_row}")
    dv_prio.add(f"I{start}:I{last_row}")
    ws.freeze_panes = "A4"

    # "Missing information" sheet: what was not found in any source (skill context-research).
    missing_items = data.get("missing_info", []) or []
    if missing_items:
        wf = wb.create_sheet("Missing information")
        wf.merge_cells("A1:F1")
        wf["A1"] = f"Missing information — {module}"
        wf["A1"].fill = PatternFill("solid", fgColor="9C5700")
        wf["A1"].font = Font(name="Arial", color="FFFFFF", bold=True, size=13)
        wf["A1"].alignment = Alignment(horizontal="center", vertical="center")
        wf.row_dimensions[1].height = 24
        for i, c in enumerate(MI_COLS):
            cell = wf.cell(row=2, column=i + 1, value=c)
            cell.fill = blue
            cell.font = white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
            wf.column_dimensions[chr(ord("A") + i)].width = MI_W[i] + 6
        for r, f in enumerate(missing_items):
            for c, key in enumerate(MI_KEYS):
                cell = wf.cell(row=3 + r, column=c + 1, value=f.get(key, "Open" if key == "status" else ""))
                cell.font = data_font
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            wf.row_dimensions[3 + r].height = 60
        affected = sum(1 for c in cases if c.get("_mi"))
        note = ws.cell(row=start + len(cases) + 2, column=1,
                       value=f"Highlighted cases: they depend on missing information ({affected}). See the \"Missing information\" sheet.")
        note.font = Font(name="Arial", size=9, italic=True, color="9C5700")
    wb.save(out)


def _wrap(text, width):
    """Wraps the text to 'width' characters, keeping the existing line breaks."""
    lines = []
    for line in str(text).split("\n"):
        if line.strip() == "":
            lines.append("")
        else:
            lines.extend(textwrap.wrap(line, width=width) or [""])
    return "\n".join(lines)


def to_md(data, out):
    _need("tabulate")
    from tabulate import tabulate
    cases = data.get("cases", [])
    title = (f"Test cases — {data.get('module', '')}").strip(" —")
    parts = [f"# {title}", "", "## Summary", ""]
    rows_s = [[_wrap(c.get(k, ""), w) for k, w in zip(SUMMARY_KEYS, SUMMARY_W)] for c in cases]
    cols_s = SUMMARY_COLS
    if any(c.get("_mi") for c in cases):
        cols_s = SUMMARY_COLS + ["Missing info"]
        rows_s = [f + [_wrap(c.get("_mi", ""), 10)] for f, c in zip(rows_s, cases)]
    parts += ["```", tabulate(rows_s, headers=cols_s, tablefmt="grid"), "```", ""]
    if data.get("_total_time_txt"):
        parts += [f"**Estimated manual run time:** {data['_total_time_txt']} in total "
                  f"(formula in `templates/time-estimation.md`).", ""]
    parts += ["## Case details", ""]
    rows = [[c.get(k, "") for k in MD_KEYS] for c in cases]
    table = tabulate(rows, headers=MD_COLS, tablefmt="grid", maxcolwidths=MD_MAXW)
    parts += ["```", table, "```", ""]
    missing_items = data.get("missing_info", []) or []
    if missing_items:
        affected = [f"{c.get('id', '')} ({c['_mi']})" for c in cases if c.get("_mi")]
        parts += ["## Missing information", "",
                  "Data that does not appear in any of the sources consulted. The cases that depend on it "
                  "carry the `@missing-info` tag" + (f": {', '.join(affected)}." if affected else "."), ""]
        rows_mi = [[f.get(k, "Open" if k == "status" else "") for k in MI_KEYS] for f in missing_items]
        parts += ["```", tabulate(rows_mi, headers=MI_COLS, tablefmt="grid", maxcolwidths=MI_W), "```", ""]
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(parts) + "\n")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    flags = [a for a in sys.argv[1:] if a.startswith("--")]
    if len(args) < 2:
        print("Usage: python scripts/generate_cases.py <input.json> <output.xlsx> [--clean]")
        sys.exit(1)
    source, target = args[0], args[1]
    if not target.lower().endswith(".xlsx"):
        print("The output must end in .xlsx")
        sys.exit(1)
    with open(source, encoding="utf-8") as f:
        data = json.load(f)
    # Sort the cases by priority: the most critical first, low priority last.
    # The sort is stable: within the same priority, the original order is kept.
    data["cases"] = sorted(
        data.get("cases", []),
        key=lambda c: PRIORITY_ORDER.get(str(c.get("priority", "")).strip().lower(), 99))
    prepare_missing_info(data)
    prepare_times(data)
    to_xlsx(data, target)
    print(f"OK: {target}")
    md = re.sub(r"\.xlsx$", ".md", target, flags=re.IGNORECASE)
    to_md(data, md)
    print(f"OK: {md}")
    # With --clean, deletes the input JSON (cross-platform, without relying on 'rm').
    if "--clean" in flags:
        try:
            os.remove(source)
            print(f"Deleted temporary file: {source}")
        except OSError:
            pass


if __name__ == "__main__":
    main()
